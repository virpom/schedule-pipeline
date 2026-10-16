"""Тесты парсера .docx (на синтетическом файле нужной структуры)."""
import datetime
from io import BytesIO

from docx import Document

from parsers.polar import parse_docx


def _make_docx(rows) -> bytes:
    doc = Document()
    table = doc.add_table(rows=0, cols=4)
    for cells in rows:
        row = table.add_row().cells
        for i, val in enumerate(cells):
            row[i].text = val
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


DATE = datetime.date(2026, 10, 16)


def test_parses_basic_lesson():
    data = _make_docx([
        ["Группа", "Пара", "Предмет", "Каб."],
        ["ДП-26", "1", "Математика - Иванов И.И.", "201"],
    ])
    lessons, no_lessons = parse_docx(data, DATE)
    assert len(lessons) == 1
    lesson = lessons[0]
    assert lesson["group"] == "ДП-26"
    assert lesson["para"] == 1
    assert lesson["subject"] == "Математика"
    assert lesson["teacher"] == "Иванов И.И."
    assert lesson["room"] == "201"


def test_parses_double_para_range():
    data = _make_docx([
        ["ДП-26", "1,2", "Организация документооборота — Харькова Д.А.", "317"],
    ])
    lessons, _ = parse_docx(data, DATE)
    assert lessons[0]["para"] == 1
    assert lessons[0]["para_end"] == 2


def test_group_carries_to_next_rows():
    data = _make_docx([
        ["ДП-26", "1", "Математика - Иванов И.И.", "201"],
        ["", "2", "Физика - Петров П.П.", "305"],
    ])
    lessons, _ = parse_docx(data, DATE)
    assert len(lessons) == 2
    assert all(l["group"] == "ДП-26" for l in lessons)


def test_parses_no_lessons_marker():
    data = _make_docx([
        ["ДП-26", "", "Занятий нет. Смотреть расписание на 17.10.2026", ""],
    ])
    lessons, no_lessons = parse_docx(data, DATE)
    assert lessons == []
    assert len(no_lessons) == 1
    assert "занятий нет" in no_lessons[0]["note"].lower()


def test_skips_header_row():
    data = _make_docx([
        ["Группа", "Пара", "Предмет", "Каб."],
        ["ДП-26", "1", "Математика - Иванов И.И.", "201"],
    ])
    lessons, _ = parse_docx(data, DATE)
    assert len(lessons) == 1


def test_teacher_with_subgroup_suffix():
    data = _make_docx([
        ["ДП-26", "3", "Физкультура - Ляшко Д.В., 1-ая подгруппа", "спортзал"],
    ])
    lessons, _ = parse_docx(data, DATE)
    assert lessons[0]["teacher"] == "Ляшко Д.В., 1-ая подгруппа"


def test_empty_document_returns_empty():
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    lessons, no_lessons = parse_docx(buf.getvalue(), DATE)
    assert lessons == [] and no_lessons == []


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"  ok  {name}")
    print("OK")
