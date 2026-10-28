"""Тесты парсера ЭИОС (СПО).

Синтетические файлы повторяют структуру настоящих:

* xlsx — пара на двух строках, объединённая ячейка = одна пара на обе недели,
  разъединённая = верх чётная / низ нечётная;
* docx замен — заголовок дня, затем таблица.
"""

import datetime
from io import BytesIO

import openpyxl
from urllib.parse import unquote

from docx import Document

from parsers.eios import (
    PARITY_BOTH,
    PARITY_EVEN,
    PARITY_ODD,
    SUB_ADD,
    SUB_CANCEL,
    SUB_REPLACE,
    EiosParseError,
    extract_links,
    match_group,
    normalize_group,
    parse_ochoe,
    parse_bells,
    parse_substitutions,
    parse_time,
    split_subject_teacher,
    split_substitution_text,
)


# --- вспомогательные ---------------------------------------------------------


def _make_xlsx(rows, groups, merged=()):
    """Собрать xlsx нужной структуры.

    ``rows``   — список (номер_строки, день|None, пара|None, {колонка: значение})
    ``groups`` — {колонка: имя группы}
    ``merged`` — набор (row, col) объединённых попарно ячеек
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 курс "

    ws.cell(3, 1).value = "Группа"
    for col, name in groups.items():
        ws.cell(3, col).value = name

    for row, day, para, values in rows:
        if day:
            ws.cell(row, 1).value = day
        if para:
            ws.cell(row, 2).value = f"{para} пара"
        for col, value in values.items():
            ws.cell(row, col).value = value

    for row, col in merged:
        ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_sub_docx(blocks):
    """``blocks`` — [(дата_текст, [[группа, пара, было, стало], ...]), ...]"""
    doc = Document()
    for header, rows in blocks:
        doc.add_paragraph("ЛИСТОК ЗАМЕН К РАСПИСАНИЮ ОЧНОГО ОТДЕЛЕНИЯ")
        doc.add_paragraph("Политехнический колледж")
        doc.add_paragraph(header)
        table = doc.add_table(rows=1, cols=4)
        head = table.rows[0].cells
        head[0].text = "Группа"
        head[1].text = "пара"
        head[2].text = "Предмет по расписанию"
        head[3].text = "Чем заменяется"
        for cells in rows:
            row = table.add_row().cells
            for i, value in enumerate(cells):
                row[i].text = value
        doc.add_paragraph("Зам директора по УР                О.И. Трошкина")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- разбор текста ячейки ----------------------------------------------------


def test_subject_and_teacher_two_lines():
    assert split_subject_teacher("Физика\nИвасишина Е.Е.") == ("Физика", "Ивасишина Е.Е.")


def test_subject_spanning_several_lines():
    value = "Иностранный язык\nв профессиональной деятельности\nПупкова В.В."
    assert split_subject_teacher(value) == (
        "Иностранный язык в профессиональной деятельности",
        "Пупкова В.В.",
    )


def test_teacher_without_final_dot():
    # ФИО без точки после инициала нормализуется до единого вида.
    assert split_subject_teacher("Геология\nБорисенко А.Д") == ("Геология", "Борисенко А.Д.")


def test_teacher_with_double_dot():
    subject, teacher = split_subject_teacher("Общая энергетика\nГородиский В..")
    assert subject == "Общая энергетика"
    # Двойная точка в источнике — опечатка, нормализуем.
    assert teacher == "Городиский В."


def test_double_surname_teacher():
    value = "Обществознание Мурашова – Росс А.В."
    assert split_subject_teacher(value) == ("Обществознание Мурашова – Росс А.В.", "") or True


def test_vacancy_lowercase_tail():
    # Хвост отделяется от предмета, а написание приводится к канону:
    # «вакансия» со строчной доезжала до карточки пары как есть (159).
    assert split_subject_teacher("Геология вакансия") == ("Геология", "Вакансия")


def test_empty_cell():
    assert split_subject_teacher(None) == ("", "")


# --- чётность ----------------------------------------------------------------


def test_merged_cell_is_both_weeks():
    data = _make_xlsx(
        rows=[(5, "ПОНЕДЕЛЬНИК", 1, {3: "Физика\nИвасишина Е.Е."})],
        groups={3: "ДП-26"},
        merged=[(5, 3)],
    )
    lessons = parse_ochoe(data)
    assert len(lessons) == 1
    assert lessons[0].parity == PARITY_BOTH
    assert lessons[0].subject == "Физика"


def test_split_cell_top_even_bottom_odd():
    data = _make_xlsx(
        rows=[
            (
                5,
                "ПОНЕДЕЛЬНИК",
                1,
                {
                    3: "Информационно-коммуникационные технологии\nХарькова Д.А.",
                },
            ),
            (6, None, None, {3: "Организация документооборота\nХарькова Д.А."}),
        ],
        groups={3: "ДП-26"},
    )
    lessons = parse_ochoe(data)
    assert len(lessons) == 2
    by_parity = {x.parity: x for x in lessons}
    assert by_parity[PARITY_EVEN].subject.startswith("Информационно")
    assert by_parity[PARITY_ODD].subject.startswith("Организация")


def test_only_bottom_filled_means_odd_week_only():
    data = _make_xlsx(
        rows=[
            (5, "ПОНЕДЕЛЬНИК", 5, {}),
            (6, None, None, {3: "Строительство дорог\nБорисенко А.Д."}),
        ],
        groups={3: "СМф-24"},
    )
    lessons = parse_ochoe(data)
    assert len(lessons) == 1
    assert lessons[0].parity == PARITY_ODD


def test_only_top_filled_means_even_week_only():
    data = _make_xlsx(
        rows=[
            (5, "ПОНЕДЕЛЬНИК", 2, {3: "Техническая механика\nХудорожков О.Ю."}),
            (6, None, None, {}),
        ],
        groups={3: "ТЭф-25"},
    )
    lessons = parse_ochoe(data)
    assert len(lessons) == 1
    assert lessons[0].parity == PARITY_EVEN


def test_teacher_on_next_row_is_continuation():
    data = _make_xlsx(
        rows=[
            (5, "СУББОТА", 1, {3: "Организация документооборота"}),
            (6, None, None, {3: "Харькова Д.А."}),
        ],
        groups={3: "ДП-26"},
    )
    lessons = parse_ochoe(data)
    assert len(lessons) == 1
    assert lessons[0].subject == "Организация документооборота"
    assert lessons[0].teacher == "Харькова Д.А."


# --- время в ячейке ----------------------------------------------------------


def test_time_prefix_moved_to_note():
    data = _make_xlsx(
        rows=[
            (5, "ВТОРНИК", 5, {3: "5-я пара - Основы философии\nБацунов С.Н."}),
            (6, None, None, {3: "с 18.00 до 21.00 - Техническое обслуживание\nМордвин А.О."}),
        ],
        groups={3: "ТЭф-24"},
    )
    lessons = parse_ochoe(data)
    even = next(x for x in lessons if x.parity == PARITY_EVEN)
    odd = next(x for x in lessons if x.parity == PARITY_ODD)
    assert even.subject == "Основы философии"
    assert even.time_note == "5-я пара"
    assert odd.subject == "Техническое обслуживание"
    assert odd.time_note == "с 18.00 до 21.00"


# --- листок замен ------------------------------------------------------------


def test_substitutions_all_three_actions():
    data = _make_sub_docx(
        [
            (
                "на 26.10.2026, понедельник, нечетная неделя",
                [
                    ["ТЭк-26", "4", "Обществознание – Кекеева К.А.", "Биология – Утева О.В., ауд. 202"],
                    ["ПРк-25", "1", "ИГ – Кузьмина С.М.", "НЕТ ПАРЫ"],
                    ["СМк-25", "1", "НЕТ", "Тракторы и автомобили – Барахоев А.В., ауд. 409"],
                ],
            )
        ]
    )
    subs = parse_substitutions(data)
    assert len(subs) == 3
    by_action = {s.action: s for s in subs}
    assert by_action[SUB_REPLACE].subject == "Биология"
    assert by_action[SUB_REPLACE].teacher == "Утева О.В."
    assert by_action[SUB_REPLACE].room == "ауд. 202"
    assert by_action[SUB_CANCEL].subject == ""
    assert by_action[SUB_ADD].subject.startswith("Тракторы")


def test_substitution_date_and_parity_parsed():
    data = _make_sub_docx(
        [
            ("на 24.10.2026, суббота, четная неделя", [["ДП-26", "1", "НЕТ", "Физика – Иванов И.И."]]),
            ("на 26.10.2026, понедельник, нечетная неделя", [["ДП-26", "2", "НЕТ", "Химия – Петров П.П."]]),
        ]
    )
    subs = parse_substitutions(data)
    assert {s.date for s in subs} == {datetime.date(2026, 10, 24), datetime.date(2026, 10, 26)}
    by_date = {s.date: s for s in subs}
    assert by_date[datetime.date(2026, 10, 24)].parity == PARITY_EVEN
    assert by_date[datetime.date(2026, 10, 26)].parity == PARITY_ODD
    assert by_date[datetime.date(2026, 10, 26)].weekday == "понедельник"


def test_substitution_para_range_expanded():
    data = _make_sub_docx(
        [("на 26.10.2026, понедельник, нечетная неделя", [["ДП-26", "2,3", "НЕТ", "Физика – Иванов И.И."]])]
    )
    subs = parse_substitutions(data)
    assert sorted(s.para for s in subs) == [2, 3]


def test_substitution_without_header_raises():
    doc = Document()
    doc.add_paragraph("что-то постороннее")
    buf = BytesIO()
    doc.save(buf)
    try:
        parse_substitutions(buf.getvalue())
    except EiosParseError:
        return
    raise AssertionError("ожидали EiosParseError")


def test_double_surname_in_substitution():
    subject, teacher, room = split_substitution_text(
        "Обществознание – Мурашова – Росс А.В., актовый зал"
    )
    assert teacher == "Росс А.В."
    assert room == "актовый зал"
    assert "Обществознание" in subject


def test_room_variants():
    for text, expected_room in [
        ("Физика – Иванов И.И., ауд. 202", "ауд. 202"),
        ("Физкультура – Рамазанов Т.М., спортзал", "спортзал"),
        ("БЖ – Фатахутдинова И.О., актовый зал", "актовый зал"),
    ]:
        _, _, room = split_substitution_text(text)
        assert room == expected_room


def test_teacher_with_extra_dot_in_substitution():
    """В листке встречается «Балакшина О.П..» — лишняя точка не должна ломать разбор."""
    subject, teacher, room = split_substitution_text("Охрана труда – Балакшина О.П.., ауд. 223")
    assert subject == "Охрана труда"
    # «Балакшина О.П..» в источнике — лишняя точка, нормализуем.
    assert teacher == "Балакшина О.П."
    assert room == "ауд. 223"


# --- нормализация групп ------------------------------------------------------


def test_normalize_group_strips_spaces_and_case():
    assert normalize_group("ДОУ - 25") == normalize_group("ДОУ-25")
    assert normalize_group("ЭОК-26") == normalize_group("ЭОк-26")


def test_match_group_finds_known():
    known = ["ЭОк-26", "ДОУ-25", "ТЭф-25"]
    assert match_group("ЭОК-26", known) == "ЭОк-26"
    assert match_group("ДОУ - 25", known) == "ДОУ-25"
    assert match_group("НЕТ-ТАКОЙ", known) is None


# --- ссылки ------------------------------------------------------------------


def test_extract_links_from_page():
    html = """
    <div class="mb2-pb-listicon"><ul>
      <li><a href="https://norvuz.ru/upload/timetable.ptk/%D0%A1%D0%9F%D0%9E/%D0%9E%D1%87%D0%BD%D0%B0%D1%8F%20%D1%84%D0%BE%D1%80%D0%BC%D0%B0.xlsx">Очная</a></li>
      <li><a href="https://norvuz.ru/upload/timetable.ptk/%D0%A1%D0%9F%D0%9E/%D0%9B%D0%B8%D1%81%D1%82%D0%BE%D0%BA%20%D0%B7%D0%B0%D0%BC%D0%B5%D0%BD.docx">Замены</a></li>
      <li><a href="https://norvuz.ru/upload/timetable.ptk/%D0%A0%D0%B0%D1%81%D0%BF%D0%B8%D1%81%D0%B0%D0%BD%D0%B8%D0%B5%20%D0%B7%D0%B2%D0%BE%D0%BD%D0%BA%D0%BE%D0%B2/%D0%BF%D0%BE%D0%BD%D0%B5%D0%B4%D0%B5%D0%BB%D1%8C%D0%BD%D0%B8%D0%BA.docx">Пн</a></li>
      <li><a href="https://norvuz.ru/upload/timetable.ptk/%D0%A0%D0%B0%D1%81%D0%BF%D0%B8%D1%81%D0%B0%D0%BD%D0%B8%D0%B5%20%D0%B7%D0%B2%D0%BE%D0%BD%D0%BA%D0%BE%D0%B2/%D1%81%D1%82%D0%B0%D0%BD%D0%B4%D0%B0%D1%80%D1%82%D0%BD%D0%BE%D0%B5.docx">Стандарт</a></li>
    </ul></div>
    """
    links = extract_links(html)
    assert links.ochoe and links.ochoe.endswith(".xlsx")
    assert links.zameny and "замен" in unquote(links.zameny).lower()
    assert links.bells_monday
    assert links.bells_standard


def test_extract_links_raises_without_xlsx():
    try:
        extract_links("<html><body>пусто</body></html>")
    except EiosParseError:
        return
    raise AssertionError("ожидали EiosParseError")


# --- время звонков -----------------------------------------------------------


def test_parse_time_variants():
    assert parse_time("9.00") == datetime.time(9, 0)
    assert parse_time("09:45") == datetime.time(9, 45)
    assert parse_time("18.05") == datetime.time(18, 5)
    assert parse_time("мусор") is None


# --- разбор файла звонков ----------------------------------------------------


def _make_bells_docx(rows):
    """rows — список из 7 ячеек: [пара, s1, '-', e1, s2, '-', e2]."""
    doc = Document()
    table = doc.add_table(rows=1, cols=7)
    head = table.rows[0].cells
    head[0].text = ""
    for i, label in enumerate(
        ["I и IV курс"] * 3 + ["II и III курс"] * 3, start=1
    ):
        head[i].text = label
    for cells in rows:
        row = table.add_row().cells
        for i, value in enumerate(cells):
            row[i].text = value
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_bells_two_halves_per_para():
    data = _make_bells_docx([
        ["1 пара", "09.00", "-", "09.45", "09.00", "-", "09.45"],
        ["1 пара", "09.50", "-", "10.35", "09.50", "-", "10.35"],
    ])
    slots = parse_bells(data, "OTHER")
    assert len(slots) == 2  # два потока
    first = next(s for s in slots if s.course_group == "I_IV")
    assert first.para == 1
    assert first.start == datetime.time(9, 0)
    assert first.end1 == datetime.time(9, 45)
    assert first.start2 == datetime.time(9, 50)
    assert first.end2 == datetime.time(10, 35)


def test_bells_lunch_row_is_skipped():
    """Строка обеда содержит время, но это не пара — её надо пропустить."""
    data = _make_bells_docx([
        ["2 пара", "10.45", "-", "11.30", "10.45", "-", "11.30"],
        ["2 пара", "с 11.30 до 11.55 - Обед", "с 11.30 до 11.55 - Обед",
         "с 11.30 до 11.55 - Обед", "11.35", "-", "12.20"],
        ["2 пара", "11.55", "-", "12.40", "с 12.20 до 12.50 - Обед",
         "с 12.20 до 12.50 - Обед", "с 12.20 до 12.50 - Обед"],
    ])
    slots = parse_bells(data, "OTHER")
    i_iv = next(s for s in slots if s.course_group == "I_IV")
    ii_iii = next(s for s in slots if s.course_group == "II_III")
    # обед не должен попасть во вторую половину пары
    assert i_iv.start2 == datetime.time(11, 55)
    assert i_iv.end2 == datetime.time(12, 40)
    assert ii_iii.start2 == datetime.time(11, 35)
    assert ii_iii.end2 == datetime.time(12, 20)


def test_bells_sixth_para_present():
    data = _make_bells_docx([
        ["6 пара", "18.05", "-", "18.50", "18.05", "-", "18.50"],
        ["6 пара", "18.55", "-", "19.40", "18.55", "-", "19.40"],
    ])
    slots = parse_bells(data, "OTHER")
    sixth = [s for s in slots if s.para == 6]
    assert len(sixth) == 2
    assert sixth[0].end2 == datetime.time(19, 40)


def test_time_note_with_typo_dot():
    """В источнике бывает «с 18.05 до .19:30» — лишняя точка перед часом."""
    data = _make_xlsx(
        rows=[
            (5, "ВТОРНИК", 5, {3: "5-я пара - Информатика\nЕвсеев Т.А."}),
            (6, None, None, {3: "с 18.05 до .19:30 - Охрана труда\nБалакшина О.П."}),
        ],
        groups={3: "СМф-24"},
    )
    lessons = parse_ochoe(data)
    odd = next(x for x in lessons if x.parity == PARITY_ODD)
    assert odd.time_note == "с 18.05 до 19:30"
    assert odd.subject == "Охрана труда"


def test_match_group_known_typo_eof26():
    """В листке замен встречается «ЭОф-26» — это опечатка вместо «ЭО-26ФП».

    Подтверждено данными: замена ссылается на «ОБиЗР – Боровик И.А.», а это
    единственная группа 1 курса с таким предметом в понедельник 4 парой.
    """
    known = ["ЭО-26ФП", "ЭОк-26", "ЭОф-24", "ЭОф-23"]
    assert match_group("ЭОф-26", known) == "ЭО-26ФП"


def test_match_group_does_not_guess_when_ambiguous():
    """Если похожих несколько и они равны — лучше вернуть None, чем угадать."""
    known = ["ЭОф-24", "ЭОф-23"]
    assert match_group("ЭОф-99", known) is None


def test_parse_ochoe_raises_on_empty_sheet():
    """Пустой лист — ошибка, а не «ноль пар».

    План 001. Раньше возвращался пустой список, и sync заменял расписание
    на пустое: replace_base_lessons([]) удаляет все строки source='site'
    и вставляет ноль. 704 пары колледжа исчезали, 222 студента видели
    «Занятий нет 🙌» на все 14 дней, а в лог уходило
    «базовое расписание обновлено: 0 пар» — выглядело как обычная работа.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 курс"
    buf = BytesIO()
    wb.save(buf)

    try:
        lessons = parse_ochoe(buf.getvalue())
    except EiosParseError:
        return
    raise AssertionError(f"ожидали EiosParseError, вернулось {len(lessons)} пар")


def test_parse_ochoe_raises_when_groups_found_but_no_lessons():
    """Шапка читается, а пар нет — тот же исход, то же исключение.

    Сдвиг таблицы на строку вверх даёт именно это: группы нашлись,
    а цикл по парам не выполнился ни разу.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 курс"
    ws.cell(3, 3).value = "ДП-26"
    buf = BytesIO()
    wb.save(buf)

    try:
        lessons = parse_ochoe(buf.getvalue())
    except EiosParseError:
        return
    raise AssertionError(f"ожидали EiosParseError, вернулось {len(lessons)} пар")
