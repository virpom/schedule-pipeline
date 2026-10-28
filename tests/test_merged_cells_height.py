"""Объединение ячейки любой высоты читается как «одна пара на обе недели» (план 102).

Условие было `max_row == min_row + 1` — ровно две строки. В файле
встречаются объединения высотой 3 и 10-11, и если методист склеит три
строки под пару, парсер молча теряет данные: ячейка уходит в обычную
ветку, верх читается как чётная неделя, низ — как нечётная, и пара
двоится вместо одной.

На текущем файле потерь нет (704 пары до и после правки), поэтому тест
на синтетической книге: та же пара, объединённая на 2 и на 3 строки,
должна дать одинаковый результат — одну пару BOTH.
"""
import io

import openpyxl

from parsers.eios import parse_ochoe

DAYS = ["ПОНЕДЕЛЬНИК"]


def _book(merge_rows: int) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 курс"
    ws["C3"] = "ДП-26"
    ws["A5"] = "ПОНЕДЕЛЬНИК"
    ws["B5"] = "1 пара"
    # «Физика» + преподаватель через перевод строки — формат колонки группы
    ws["C5"] = "Физика\nИванов И.И."
    if merge_rows > 1:
        ws.merge_cells(start_row=5, start_column=3,
                       end_row=5 + merge_rows - 1, end_column=3)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _paras(data: bytes):
    return [l for l in parse_ochoe(data) if l.group == "ДП-26"]


def test_two_row_merge_is_both():
    got = _paras(_book(2))
    assert len(got) == 1, got
    assert got[0].parity == "BOTH"
    assert "Физика" in got[0].subject


def test_three_row_merge_is_both():
    """Объединение на три строки — тоже одна пара, а не две."""
    got = _paras(_book(3))
    assert len(got) == 1, got
    assert got[0].parity == "BOTH", got[0].parity


def test_height_does_not_change_result():
    assert len(_paras(_book(2))) == len(_paras(_book(3))) == 1
