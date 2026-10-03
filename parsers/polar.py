import datetime
import re
from io import BytesIO

from aiohttp import ClientSession
from bs4 import BeautifulSoup
from docx import Document

DAILY_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")


async def fetch_html(session: ClientSession, url: str) -> str:
    async with session.get(url, raise_for_status=True) as response:
        return await response.text()


def extract_daily_links(html: str, base_url: str) -> list[tuple[datetime.date, str]]:
    soup = BeautifulSoup(html, "html.parser")
    result = []
    for a in soup.find_all("a", href=True):
        text = a.get_text(strip=True)
        if "Расписание на" not in text:
            continue
        m = DAILY_RE.search(text) or DAILY_RE.search(a["href"])
        if not m:
            continue
        day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        url = a["href"]
        if url.startswith("/"):
            url = base_url.rstrip("/") + url
        result.append((datetime.date(year, month, day), url))
    return result


def _split_subject(raw: str) -> tuple[str, str]:
    raw = raw.strip()
    for sep in (" – ", " — ", " - "):
        if sep in raw:
            subject, _, teacher = raw.partition(sep)
            return subject.strip(), teacher.strip()
    return raw, ""


def parse_docx(data: bytes, date: datetime.date) -> list[dict]:
    doc = Document(BytesIO(data))
    if not doc.tables:
        return []

    lessons = []
    current_group = ""
    for row in doc.tables[0].rows:
        cells = [c.text.strip() for c in row.cells]
        if len(cells) < 4:
            continue
        group, para, subject_raw, room = cells[0], cells[1], cells[2], cells[3]

        if group and group.lower().startswith("группа"):
            continue
        if group:
            current_group = group
        if not para.isdigit():
            continue

        subject, teacher = _split_subject(subject_raw)
        if not subject:
            continue

        lessons.append({
            "date": date,
            "group": current_group,
            "para": int(para),
            "subject": subject,
            "teacher": teacher,
            "room": room,
        })
    return lessons


if __name__ == "__main__":
    import asyncio

    async def demo():
        async with ClientSession() as session:
            html = await fetch_html(
                session,
                "https://polaruniversity.ru/obuchayushchimsya/raspisanie-zanyatiy/",
            )
            links = extract_daily_links(html, "https://polaruniversity.ru")
            assert links, "no daily links found"
            date, url = links[0]
            async with session.get(url, raise_for_status=True) as resp:
                lessons = parse_docx(await resp.read(), date)
            assert lessons, "no lessons parsed"
            print(f"{url} -> {len(lessons)} lessons on {date}")
            for l in lessons[:3]:
                print("  ", l)

    asyncio.run(demo())
